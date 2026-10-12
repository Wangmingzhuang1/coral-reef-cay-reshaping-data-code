
import sys as _project_sys
from pathlib import Path as _ProjectPath
_project_sys.path.insert(0, str(_ProjectPath(__file__).resolve().parents[1] / "公共"))
from 项目路径 import ROOT as PROJECT_ROOT, configure_imports
configure_imports()
import re
import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Pt, RGBColor


ROOT = PROJECT_ROOT
MD_PATH = ROOT / "论文初稿.md"
DOCX_PATH = ROOT / "论文" / "中文论文.docx"
SUPPLEMENT_MD_PATH = ROOT / "论文补充材料.md"
SUPPLEMENT_DOCX_PATH = ROOT / "论文" / "中文论文补充材料.docx"
FIGURE_DIR = ROOT / "outputs" / "figures"

_INDEX = json.loads((FIGURE_DIR / "当前图件索引.json").read_text(encoding="utf-8"))
FIGURES = {int(x["图号"]): ROOT / x["文件"]["png"]["path"] for x in _INDEX["图件"] if not x["图号"].startswith("S")}

def _parse_captions(path, prefix):
    caps = {}
    for raw in path.read_text(encoding="utf-8").splitlines():
        line = raw.strip()
        if not line.startswith("**" + prefix):
            continue
        body = line.replace("**", "")
        key = body.split("｜", 1)[0].replace(prefix, "").strip()
        caps[key] = body
    return caps


FIGURE_CAPTIONS = {int(k): v for k, v in _parse_captions(MD_PATH, "图 ").items()}

SUPPLEMENTARY_FIGURES = {x["图号"]: ROOT / x["文件"]["png"]["path"] for x in _INDEX["图件"] if x["图号"].startswith("S")}

SUPPLEMENTARY_FIGURE_CAPTIONS = _parse_captions(SUPPLEMENT_MD_PATH, "补充图 ")

SUPPLEMENTARY_SECTION_FIGURES = {key:[key] for key in SUPPLEMENTARY_FIGURES}


def set_run_font(run, east_asia="宋体", latin="Times New Roman", size=11, bold=None, italic=None):
    run.font.name = latin
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), latin)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), latin)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=60, start=100, bottom=60, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D9D9D9", size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:color"), color)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char_1 = OxmlElement("w:fldChar")
    fld_char_1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char_2 = OxmlElement("w:fldChar")
    fld_char_2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char_1, instr_text, fld_char_2])
    set_run_font(run, size=9)


def clear_paragraph(paragraph):
    for child in list(paragraph._p):
        paragraph._p.remove(child)


def remove_paragraph_border(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is not None:
        p_pr.remove(borders)


def append_delta_equation(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    omath_para = OxmlElement("m:oMathPara")
    omath = OxmlElement("m:oMath")

    def math_run(text):
        r = OxmlElement("m:r")
        t = OxmlElement("m:t")
        t.text = text
        r.append(t)
        return r

    omath.append(math_run("ΔAᵢ = |"))
    frac = OxmlElement("m:f")
    num = OxmlElement("m:num")
    num.append(math_run("Aᵢ,late − Aᵢ,early"))
    den = OxmlElement("m:den")
    den.append(math_run("Aᵢ,early"))
    frac.extend([num, den])
    omath.append(frac)
    omath.append(math_run("|"))
    omath_para.append(omath)
    paragraph._p.append(omath_para)


def append_ctmc_equation(paragraph):
    """写入三状态连续时间模型的观测转移概率公式。"""
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    omath_para = OxmlElement("m:oMathPara")
    omath = OxmlElement("m:oMath")
    math_run = OxmlElement("m:r")
    math_text = OxmlElement("m:t")
    math_text.text = "Pᵢⱼ(Δt) = [exp(QΔt)]ᵢⱼ"
    math_run.append(math_text)
    omath.append(math_run)
    omath_para.append(omath)
    paragraph._p.append(omath_para)


def clean_inline(text):
    leading = " " if text.startswith(" ") else ""
    trailing = " " if text.endswith(" ") else ""
    replacements = {
        r"\(\beta": "β",
        r"\rho": "ρ",
        r"\Delta": "Δ",
        r"\geq": "≥",
        r"\leq": "≤",
        r"\times": "×",
        r"^\circ": "°",
        r"\circ": "°",
        r"\mathrm C": "C",
        r"\,": " ",
        r"\ ": " ",
    }
    for old, new in replacements.items():
        text = text.replace(old, new)
    text = text.replace(r"\(", "").replace(r"\)", "")
    text = text.replace("$", "")
    text = re.sub(r"\^\{?(-?\d+)\}?", lambda m: str(m.group(1)).translate(str.maketrans("0123456789-", "⁰¹²³⁴⁵⁶⁷⁸⁹⁻")), text)
    text = text.replace(r"\mathrm{m\,s^{-1}}", "m s⁻¹")
    text = text.replace(r"\mathrm", "")
    # 仅转换独立数学变量的下标，保留GBR_101、sand_cay_id等数据身份。
    subscripts = {"ii": "ᵢᵢ", "ij": "ᵢⱼ", "i": "ᵢ", "1": "₁"}
    text = re.sub(r"(?<![\w])([A-Za-z])_(ii|ij|i|1)(?![\w])",
                  lambda m: m.group(1) + subscripts[m.group(2)], text)
    text = text.replace("{", "").replace("}", "")
    text = text.replace("`", "")
    text = text.replace("**", "")
    text = re.sub(r"</?sup>", "", text)
    return leading + text.strip() + trailing


def add_rich_paragraph(doc, text, style=None, first_indent=True):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(4)
    if first_indent and style is None:
        p.paragraph_format.first_line_indent = Pt(22)

    parts = re.split(r"(<sup>[^<]+</sup>|\*\*[^*]+\*\*|\*[^*]+\*)", text)
    for part in parts:
        if not part:
            continue
        superscript = part.startswith("<sup>") and part.endswith("</sup>")
        bold = part.startswith("**") and part.endswith("**")
        italic = not bold and part.startswith("*") and part.endswith("*")
        if superscript:
            value = clean_inline(part[5:-6])
        elif bold:
            value = clean_inline(part[2:-2])
        elif italic:
            value = clean_inline(part[1:-1])
        else:
            value = clean_inline(part)
        if value:
            run = p.add_run(value)
            set_run_font(run, bold=True if bold else None, italic=True if italic else None)
            if superscript:
                run.font.superscript = True
    return p


def add_author_information(doc, content):
    p = add_rich_paragraph(doc, content, first_indent=False)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.line_spacing = 1.1
    p.paragraph_format.space_after = Pt(6)
    for run in p.runs:
        set_run_font(run, size=10 if content.startswith(("<sup>", "* Corresponding", "Email addresses:")) else 11)
    return p


def is_author_information(content):
    return content.startswith(("Mingzhuang Wang<sup>", "<sup>", "* Corresponding authors:", "Email addresses:"))


def parse_markdown(path=MD_PATH):
    lines = path.read_text(encoding="utf-8").splitlines()
    blocks = []
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line:
            i += 1
            continue
        if line.startswith("\\["):
            eq = [line]
            while not line.endswith("\\]") and i + 1 < len(lines):
                i += 1
                line = lines[i].rstrip()
                eq.append(line)
            blocks.append(("equation", "\n".join(eq)))
            i += 1
            continue
        if line.startswith("|") and i + 1 < len(lines) and re.match(r"^\|\s*:?-{3,}", lines[i + 1].lstrip()):
            rows = [line]
            i += 2
            while i < len(lines) and lines[i].startswith("|"):
                rows.append(lines[i].rstrip())
                i += 1
            parsed = [[clean_inline(c.strip()) for c in row.strip("|").split("|")] for row in rows]
            blocks.append(("table", parsed))
            continue
        if line.startswith("# "):
            blocks.append(("title", clean_inline(line[2:])))
            i += 1
            continue
        if line.startswith("## "):
            blocks.append(("h1", clean_inline(line[3:])))
            i += 1
            continue
        if line.startswith("### "):
            blocks.append(("h2", clean_inline(line[4:])))
            i += 1
            continue
        para = [line]
        while i + 1 < len(lines) and lines[i + 1].strip() and not lines[i + 1].startswith(("#", "|", "\\[")):
            i += 1
            para.append(lines[i].strip())
        blocks.append(("paragraph", " ".join(para)))
        i += 1
    return blocks


def clear_body_preserve_section(doc):
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def add_table(doc, rows):
    table = doc.add_table(rows=1, cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_table_borders(table)
    header_pr = table.rows[0]._tr.get_or_add_trPr()
    header_pr.append(OxmlElement("w:tblHeader"))
    for j, value in enumerate(rows[0]):
        cell = table.rows[0].cells[j]
        cell.text = value
        set_cell_shading(cell, "D9EAF7")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell)
        cell.paragraphs[0].paragraph_format.keep_with_next = True
        for run in cell.paragraphs[0].runs:
            set_run_font(run, size=9, bold=True)
    for row_index, values in enumerate(rows[1:], start=1):
        cells = table.add_row().cells
        for j, value in enumerate(values):
            cell = cells[j]
            cell.text = value
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if row_index % 2 == 0:
                set_cell_shading(cell, "F7FAFC")
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                for run in p.runs:
                    set_run_font(run, size=8.2)
    for row in table.rows:
        row_pr = row._tr.get_or_add_trPr()
        if row_pr.find(qn("w:cantSplit")) is None:
            row_pr.append(OxmlElement("w:cantSplit"))
    # 九列表的来源与样本列较窄，为带符号系数和完整小数留出明确宽度。
    if len(rows[0]) == 9 and "双聚类CI下限" in rows[0]:
        table.autofit = False
        widths = [35, 32, 36, 36, 56, 60, 60, 50, 50]
        for j, width in enumerate(widths):
            table.columns[j].width = Pt(width)
            for cell in table.columns[j].cells:
                cell.width = Pt(width)
                set_cell_margins(cell, start=40, end=40)
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.first_line_indent = Pt(0)
                    paragraph.paragraph_format.left_indent = Pt(0)
                    paragraph.paragraph_format.right_indent = Pt(0)
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if len(table.rows) <= 6:
        for row in table.rows[:-1]:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.keep_with_next = True
    doc.add_paragraph()


def add_figure(doc, number):
    path = FIGURES[number]
    if not path.exists():
        raise FileNotFoundError(f"主图缺失，停止文档生成：{path}")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    if number == 1:
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_after = Pt(2)
    run = p.add_run()
    run.add_picture(str(path), width=Cm(15.4))
    caption = doc.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_after = Pt(8)
    caption.paragraph_format.keep_together = True
    if number == 1:
        caption.paragraph_format.line_spacing = 1.05
    run = caption.add_run(FIGURE_CAPTIONS[number])
    set_run_font(run, size=9)


def add_supplementary_figure(doc, number):
    path = SUPPLEMENTARY_FIGURES[number]
    if not path.exists():
        raise FileNotFoundError(f"补充图缺失，停止文档生成：{path}")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    run = p.add_run()
    run.add_picture(str(path), width=Cm(15.4))
    caption = doc.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_after = Pt(8)
    caption.paragraph_format.keep_together = True
    run = caption.add_run(SUPPLEMENTARY_FIGURE_CAPTIONS[number])
    set_run_font(run, size=9)


def main():
    doc = Document(DOCX_PATH)
    clear_body_preserve_section(doc)
    section = doc.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)
    section.header_distance = Cm(1.2)
    section.footer_distance = Cm(1.2)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(4)

    title = styles["Title"]
    title.font.name = "Times New Roman"
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    title.font.size = Pt(16)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title_borders = title._element.pPr.find(qn("w:pBdr"))
    if title_borders is not None:
        title._element.pPr.remove(title_borders)

    for name, size in (("Heading 1", 14), ("Heading 2", 12)):
        style = styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True

    blocks = parse_markdown()
    inserted = set()
    pending_figures = []
    skip_figure_legend_paragraphs = False
    skip_internal_audit = False
    in_references = False
    in_results = False
    for kind, content in blocks:
        # 完成当前小节的正文后再插图，保持发现与证据的连续阅读。
        if kind in ("h1", "h2") and pending_figures:
            for number in pending_figures:
                add_figure(doc, number)
                inserted.add(number)
            pending_figures.clear()
        if kind == "h1" and content == "内部证据核查（非论文正文）":
            skip_internal_audit = True
            continue
        if kind == "h1" and content == "数据与代码可用性":
            skip_internal_audit = False
        if skip_internal_audit:
            continue
        if kind == "title":
            p = doc.add_paragraph(style="Title")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(16)
            remove_paragraph_border(p)
            run = p.add_run(content)
            set_run_font(run, east_asia="黑体", size=16, bold=True)
        elif kind == "h1":
            in_results = content == "3 结果"
            in_references = content == "参考文献"
            skip_figure_legend_paragraphs = content == "图注"
            if skip_figure_legend_paragraphs:
                continue
            p = doc.add_paragraph(content, style="Heading 1")
            for run in p.runs:
                set_run_font(run, east_asia="黑体", size=14, bold=True)
        elif kind == "h2":
            p = doc.add_paragraph(content, style="Heading 2")
            for run in p.runs:
                set_run_font(run, east_asia="黑体", size=12, bold=True)
        elif kind == "equation":
            p = doc.add_paragraph()
            if "P_{ij}" in content:
                append_ctmc_equation(p)
            else:
                append_delta_equation(p)
        elif kind == "table":
            add_table(doc, content)
        elif kind == "paragraph":
            if is_author_information(content):
                add_author_information(doc, content)
                continue
            if skip_figure_legend_paragraphs:
                continue
            if content.startswith("**关键词：**"):
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(8)
                run = p.add_run("关键词：")
                set_run_font(run, bold=True)
                run = p.add_run(clean_inline(content.replace("**关键词：**", "", 1)))
                set_run_font(run)
            else:
                is_reference = in_references
                p = add_rich_paragraph(doc, content, first_indent=not is_reference)
                if is_reference:
                    p.paragraph_format.left_indent = Pt(22)
                    p.paragraph_format.first_line_indent = Pt(-22)
                    p.paragraph_format.line_spacing = 1.1
                    p.paragraph_format.space_after = Pt(1)
                    for run in p.runs:
                        set_run_font(run, size=10, italic=run.italic)
            for n in FIGURES:
                if in_results and re.search(rf"(?<!补充)图\s*{n}(?!\d)", clean_inline(content)) and n not in inserted and n not in pending_figures:
                    pending_figures.append(n)

    for number in pending_figures:
        add_figure(doc, number)
        inserted.add(number)

    for paragraph in section.header.paragraphs:
        clear_paragraph(paragraph)
    for paragraph in section.footer.paragraphs:
        clear_paragraph(paragraph)
    add_page_number(section.footer.paragraphs[0])
    doc.core_properties.title = MD_PATH.read_text(encoding="utf-8").splitlines()[0].lstrip("# ")
    doc.core_properties.subject = "中文论文初稿"
    doc.core_properties.keywords = "珊瑚礁沙洲; 动态平衡; 植被; 台风; ENSO"
    doc.save(DOCX_PATH)

    supplement = Document(DOCX_PATH)
    clear_body_preserve_section(supplement)
    supplement_section = supplement.sections[0]
    inserted_supplementary = set()
    pending_supplementary = None
    skip_supplementary_legends = False
    for kind, content in parse_markdown(SUPPLEMENT_MD_PATH):
        if kind == "title":
            p = supplement.add_paragraph(style="Title")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(16)
            remove_paragraph_border(p)
            run = p.add_run(content)
            set_run_font(run, east_asia="黑体", size=16, bold=True)
        elif kind == "h1":
            skip_supplementary_legends = content == "补充图注"
            if skip_supplementary_legends:
                continue
            p = supplement.add_paragraph(content, style="Heading 1")
            for run in p.runs:
                set_run_font(run, east_asia="黑体", size=14, bold=True)
            pending_supplementary = SUPPLEMENTARY_SECTION_FIGURES.get(content.split(" ", 1)[0])
            if isinstance(pending_supplementary, str):
                pending_supplementary = [pending_supplementary]
        elif kind == "h2":
            p = supplement.add_paragraph(content, style="Heading 2")
            for run in p.runs:
                set_run_font(run, east_asia="黑体", size=12, bold=True)
        elif kind == "table":
            add_table(supplement, content)
        elif kind == "equation":
            p = supplement.add_paragraph()
            if "Delta A" in content:
                append_delta_equation(p)
            else:
                append_ctmc_equation(p)
        elif kind == "paragraph":
            if is_author_information(content):
                add_author_information(supplement, content)
                continue
            if skip_supplementary_legends:
                continue
            add_rich_paragraph(supplement, content)
            if pending_supplementary:
                for figure_key in pending_supplementary:
                    if figure_key not in inserted_supplementary:
                        add_supplementary_figure(supplement, figure_key)
                        inserted_supplementary.add(figure_key)
                pending_supplementary = None
    for paragraph in supplement_section.header.paragraphs:
        clear_paragraph(paragraph)
    for paragraph in supplement_section.footer.paragraphs:
        clear_paragraph(paragraph)
    add_page_number(supplement_section.footer.paragraphs[0])
    supplement.core_properties.title = "珊瑚礁沙洲形态调整研究补充材料"
    supplement.core_properties.subject = "中文论文补充材料"
    supplement.core_properties.keywords = "珊瑚礁沙洲; 台风; ENSO; ERA5; 补充材料"
    supplement.save(SUPPLEMENT_DOCX_PATH)
    print(f"saved={DOCX_PATH}")
    print(f"saved={SUPPLEMENT_DOCX_PATH}")
    print(f"figures_inserted={sorted(inserted)}")
    print(f"supplementary_figures_inserted={sorted(inserted_supplementary)}")


if __name__ == "__main__":
    main()
