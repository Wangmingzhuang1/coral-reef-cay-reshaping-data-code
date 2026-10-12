"""核验实际Word图片、原件指纹及现行图号，不以登记数量代替实物。"""

import sys as _project_sys
from pathlib import Path as _ProjectPath
_project_sys.path.insert(0, str(_ProjectPath(__file__).resolve().parents[1] / "公共"))
from 项目路径 import ROOT as PROJECT_ROOT, configure_imports
configure_imports()
from pathlib import Path
import hashlib, json, re
from collections import Counter
import zipfile
import pandas as pd
import numpy as np
from docx import Document

ROOT = PROJECT_ROOT
BASE = ROOT.parent/'归档/重构前'
def digest(path): return hashlib.sha256(path.read_bytes()).hexdigest()

def reference_entries(text):
    return Counter(line.strip() for line in text.split("## 参考文献",1)[1].splitlines() if line.strip())

def verify_reference_revision(previous, current):
    ledger_path=ROOT/"outputs/引用核验/当前引用对应核验.json"
    if not ledger_path.exists():
        assert reference_entries(previous)==reference_entries(current)
        return
    ledger=json.loads(ledger_path.read_text(encoding="utf-8"))
    assert reference_entries(previous)==Counter(ledger["original_entries"])
    assert reference_entries(current)==Counter(ledger["current_entries"])
    assert not ledger["missing"]

def main():
    manifest = json.loads((ROOT/'outputs/figures/当前图件索引.json').read_text(encoding='utf-8'))
    assert [x['图号'] for x in manifest['图件']] == [str(x) for x in range(1,7)]+['S'+str(x) for x in range(1,9)]
    checked = []
    for name, prefix, count in [('中文论文.docx','',6),('中文论文补充材料.docx','S',8)]:
        document = Document(ROOT/'论文'/name)
        expected = [x for x in manifest['图件'] if x['图号'].startswith('S') == bool(prefix)]
        expected_hashes = {x['文件']['png']['sha256'] for x in expected}
        actual_hashes = set()
        for shape in document.inline_shapes:
            relation = shape._inline.graphic.graphicData.pic.blipFill.blip.embed
            actual_hashes.add(hashlib.sha256(document.part.related_parts[relation].blob).hexdigest())
        assert len(document.inline_shapes) == count, f'{name}实际图片数量错误'
        assert actual_hashes == expected_hashes, f'{name}嵌入图片与当前清单不一致'
        text='\n'.join(p.text for p in document.paragraphs)
        for x in expected:
            label = ('补充图 '+x['图号']) if prefix else ('图 '+x['图号'])
            assert label in text, f'{name}缺图注{label}'
        checked.append({'文档':name,'实际嵌入图片':len(document.inline_shapes),'图片哈希与清单一致':True,'表格数':len(document.tables)})
    for x in manifest['图件']:
        for ext, file in x['文件'].items():
            assert digest(ROOT/file['path']) == file['sha256']
    original = json.loads((BASE/'第一阶段审计/复核结果.json').read_text(encoding='utf-8'))['original_sha256']
    assert all(digest(BASE/path)==sha for path,sha in original.items())
    old=(BASE/'论文初稿.md').read_text(encoding='utf-8'); new=(ROOT/'论文初稿.md').read_text(encoding='utf-8')
    verify_reference_revision(old,new)
    for manuscript in ['论文初稿.md','论文补充材料.md']:
        text=(ROOT/manuscript).read_text(encoding='utf-8')
        assert not re.search(r'\[CITATION NEEDED\]|PLACEHOLDER_|TODO[: ]+verify',text,re.I)
        assert not re.search(r'补充图 S(?:9|1[0-2])',text), '出现已撤销图号'
        assert not re.search(r'补充图 S7[cde]',text), 'S7仅有a、b面板'
    src=ROOT/'outputs/figures/source_data'
    a=pd.read_csv(src/'表A_沙洲级面积路径状态.csv'); b=pd.read_csv(src/'表B_区间级状态边界位移环境.csv')
    assert len(a)==93 and len(b)==1490
    assert b.loc[~b.mask_geometry_comparable,'boundary_exchange_m2'].isna().all()
    assert (~b.mask_geometry_comparable).sum()==28
    enso=pd.read_csv(ROOT/'outputs/ENSO滞后模型.csv')
    supplemental=pd.read_csv(ROOT/'outputs/气候检验/ENSO总体滞后模型.csv')
    merged=enso.merge(supplemental,on=['outcome','lag_months'])
    assert len(merged)==14
    assert (abs(merged.coefficient_per_1c-merged.coefficient)<1e-10).all(), 'ENSO两套结果表系数不一致'
    audit_dir=ROOT/'outputs/轨迹与转换核查'
    continuous=pd.read_csv(audit_dir/'连续空间轨迹与首末边界交换.csv')
    endpoints=continuous[continuous.endpoint_geometry_comparable]
    assert len(continuous)==93 and len(endpoints)==90
    assert np.allclose(endpoints.endpoint_exchange_fraction,endpoints.endpoint_loss_fraction+endpoints.endpoint_gain_fraction)
    assert np.allclose(endpoints.endpoint_net_fraction,endpoints.endpoint_gain_fraction-endpoints.endpoint_loss_fraction)
    assert continuous.centroid_net_normalized.le(continuous.centroid_path_normalized+1e-8).all()
    assert continuous.centroid_net_normalized.le(2*continuous.centroid_envelope_normalized+1e-8).all()
    pair=pd.read_csv(audit_dir/'相似面积不同空间活动展示对.csv')
    assert len(pair)==2 and pair.region.nunique()==1 and pair.sensor.eq('sentinel2').all()
    for field in ['start_area_m2','span_years','n_observations']:
        assert pair[field].max()/pair[field].min()<=1.5
    assert pair.relative_change.max()-pair.relative_change.min()<=.1
    events=pd.read_csv(audit_dir/'状态转换前后候选与排除标记.csv')
    retained=events[events.non_reused]
    assert len(events)==34 and len(retained)==27 and retained.sand_cay_id.nunique()==14
    for _,group in retained.groupby('sand_cay_id'):
        dates=group[['pre_start','event_start','event_end','post_end']].to_numpy().ravel()
        assert len(dates)==len(set(dates)), '转换窗口复用同一影像日期'
    support=pd.read_csv(audit_dir/'转换前后有效样本层级.csv').set_index('specification')
    assert support.loc['strict_non_reused','n_events']==1 and support.loc['observation_matched','n_events']==0
    assert support.loc['non_reused','environment_complete']==7
    registration=pd.read_csv(audit_dir/'NH_015背景特征配准筛查.csv')
    assert len(registration)==29 and registration.screen_accepted.sum()==9
    calibration=pd.read_csv(src/'补图S5_状态模型校准.csv')
    assert calibration.loc[calibration.interval_bin.eq('180-365 d'),'observed_pairs'].iloc[0]==1
    strict=pd.read_csv(src/'补图S5_严格标签变化.csv')
    summary=json.loads((audit_dir/'证据核查.json').read_text(encoding='utf-8'))
    assert summary['NH_015']['mask_centroid_reproduction']
    shape=pd.read_csv(ROOT/'outputs/形状审查/主序列形状分布与窗口变化.csv')
    shape_source=pd.read_csv(src/'补图S8_长期形状与季节支持.csv')
    pd.testing.assert_frame_equal(shape,shape_source)
    assert len(shape)==93 and shape.shape_observations.sum()==2115 and shape.anchor_overlap.sum()==3
    assert (~shape.anchor_overlap).sum()==90
    assert shape.loc[~shape.anchor_overlap,'E_direction_preserved_1px'].sum()==53
    assert shape.loc[~shape.anchor_overlap & shape.maximum_cross_window_season_gap_days.le(45)].empty
    trace=pd.read_csv(ROOT/'outputs/形状审查/早晚掩膜来源与算法复核.csv')
    shape_audit=json.loads((ROOT/'outputs/形状审查/形状核查.json').read_text(encoding='utf-8'))
    assert shape_audit['input_sha256']==digest(ROOT/'outputs/沙洲观测主表.csv')
    assert len(trace)==462 and np.allclose(trace.stored_E,trace.recomputed_E)
    assert len(pd.read_csv(src/'补图S8_逐景轮廓来源.csv'))==10
    quarter=pd.read_csv(ROOT/'outputs/形状审查/同季度早晚形状位置.csv')
    quarter_source=pd.read_csv(src/'补图S8_同季度形状位置对照.csv')
    pd.testing.assert_frame_equal(quarter,quarter_source)
    images=pd.read_csv(ROOT/'outputs/形状审查/同季度逐景来源与边界情景.csv')
    assert len(quarter)==63 and len(images)==315 and quarter.E_direction_preserved_1px.sum()==30
    assert quarter.same_direction_as_original.sum()==38
    assert (quarter.E_direction_preserved_1px & quarter.original_E_direction_preserved_1px & quarter.same_direction_as_original).sum()==18
    assert quarter.window_separation_years.ge(3).all() and quarter.maximum_cross_window_season_gap_days.le(90).all()
    observations=pd.read_csv(ROOT/'outputs/沙洲观测主表.csv');observations.date=pd.to_datetime(observations.date)
    for _,x in quarter.iterrows():
        selected=images[(images.sand_cay_id==x.sand_cay_id)&(images.sensor==x.sensor)&(images.reference_frame_id==x.reference_frame_id)]
        assert len(selected)==5 and selected.image_id.nunique()==5
        dates=pd.to_datetime(selected.date)
        assert dates.dt.quarter.eq(x.quarter).all()
        merged=selected.merge(observations,on=['sand_cay_id','sensor','reference_frame_id','image_id'],suffixes=('_trace',''))
        assert len(merged)==5 and merged.quality_grade.isin(['A','B']).all()
        assert np.allclose(merged.E_original,merged.sand_cay_elongation)
        early=merged[merged.window.eq('early')];late=merged[merged.window.eq('late')]
        radius=np.sqrt(early.sand_cay_area_m2.median()/np.pi)
        positions=[]
        for window in [early,late]:
            positions.append(np.median(window[['sand_cay_centroid_x','sand_cay_centroid_y']].to_numpy()*window.pixel_size_m.to_numpy()[:,None],axis=0))
        assert np.isclose(np.linalg.norm(positions[1]-positions[0])/radius,x.window_centroid_net_normalized)
        assert np.isclose(late.sand_cay_elongation.median()-early.sand_cay_elongation.median(),x.delta_E)
        assert np.isclose(np.log(late.sand_cay_area_m2.median()/early.sand_cay_area_m2.median()),x.delta_log_area)
    with zipfile.ZipFile(ROOT.parent/'归档/覆盖主线与P0核查前.zip') as backup:
        prior_text=backup.read('论文初稿.md').decode('utf-8')
        verify_reference_revision(prior_text,new)
        previous=pd.read_csv(backup.open('outputs/沙洲观测主表.csv'))
        current=pd.read_csv(ROOT/'outputs/沙洲观测主表.csv')
        pd.testing.assert_frame_equal(previous,current[previous.columns],check_exact=False,rtol=1e-12,atol=1e-12)
        previous_cases=pd.read_csv(backup.open('outputs/figures/source_data/同沙洲状态图像.csv'))
        current_cases=pd.read_csv(ROOT/'outputs/figures/source_data/同沙洲状态图像.csv')
        assert previous_cases.image_id.tolist()==current_cases.image_id.tolist(), '图3绘图可以改善，原案例影像选择保持'
    p0=json.loads((ROOT/'outputs/P0证据核查/核查结果.json').read_text(encoding='utf-8'))
    assert p0['master_sha256']==digest(ROOT/'outputs/沙洲观测主表.csv')
    assert p0['continuous_track_sha256']==digest(ROOT/'outputs/轨迹与转换核查/连续空间轨迹与首末边界交换.csv')
    assert p0['code_sha256']==digest(ROOT/'代码/控制验证/核查紧致度与质心证据.py')
    assert p0['n_recomputed']==2615 and p0['n_unavailable_outside_analysis']==6
    canonical=current[current.compactness_primary_crofton.notna()]
    assert np.allclose(canonical.compactness_primary_crofton,4*np.pi*canonical.compactness_area_pixels/canonical.compactness_perimeter_crofton_pixels**2)
    assert np.allclose(canonical.compactness_area_pixels,canonical.sand_cay_area_pixels)
    annual=pd.read_csv(ROOT/'outputs/审稿问题复核/单一原始网格留一年汇总.csv');visual=pd.read_csv(src/'figure_2_harmonic_relative_error_reduction.csv')
    row=visual[visual.metric.eq('sand_cay_compactness')].iloc[0]
    strict_annual=pd.read_csv(ROOT/'outputs/审稿问题复核/单一原始网格留一年逐景对照.csv')
    means=strict_annual[strict_annual.metric.eq('sand_cay_compactness')].groupby('sand_cay_id')[['baseline_error','seasonal_error']].mean()
    assert np.isclose(row.mean_relative_error_reduction,1-means.seasonal_error.sum()/means.baseline_error.sum())
    assert np.isclose(row.fdr_q_value,annual.loc[annual.metric.eq('sand_cay_compactness'),'fdr_q_value'].iloc[0])
    paired=pd.read_csv(ROOT/'outputs/P0证据核查/全部同区域可比记录差异.csv');pd.testing.assert_frame_equal(paired,pd.read_csv(src/'主图2_全部匹配差异.csv'))
    assert len(paired)==32 and paired.non_reused_by_id.sum()==13
    non_reused=paired[paired.non_reused_by_id]
    assert len(set(non_reused.cay0)|set(non_reused.cay1))==26
    sampling=pd.read_csv(ROOT/'outputs/P0证据核查/主序列端点固定采样敏感性.csv')
    assert len(sampling)==372 and sampling.groupby('sand_cay_id').net_r.apply(lambda x:x.max()-x.min()).max()<1e-10
    registered=pd.read_csv(ROOT/'outputs/P0证据核查/核心与小尺寸案例配准支持.csv');assert len(registered)==10
    coverage=pd.read_csv(src/'主图6_覆盖与累计路径统计.csv');assert len(coverage)==2 and coverage.fdr_q_value.lt(.05).all() and coverage.spearman_rho.lt(0).all()
    from 汇总长期序列区域特征 import summarize
    regional=pd.read_csv(ROOT/'outputs/轨迹与转换核查/长期主序列区域汇总.csv')
    pd.testing.assert_frame_equal(regional,summarize(continuous),check_exact=False,rtol=1e-12,atol=1e-12)
    regional_provenance=json.loads((ROOT/'outputs/轨迹与转换核查/区域汇总来源.json').read_text(encoding='utf-8'))
    assert regional_provenance['input_sha256']==digest(ROOT/'outputs/轨迹与转换核查/连续空间轨迹与首末边界交换.csv')
    assert regional.n_primary.sum()==93 and regional_provenance['n_four_focus_regions']==91
    assert regional[['n_near_stable','n_net_gain','n_net_loss']].sum().tolist()==[63,17,13]
    assert regional.endpoint_exchange_fraction_n.sum()==90
    source_rows=[]
    for sensor,group in continuous.groupby('sensor'):
        table=summarize(group);table.insert(1,'sensor',sensor);source_rows.append(table)
    pd.testing.assert_frame_equal(pd.read_csv(ROOT/'outputs/轨迹与转换核查/长期主序列区域分源汇总.csv'),pd.concat(source_rows,ignore_index=True),check_exact=False,rtol=1e-12,atol=1e-12)
    assert '### 3.1 面积变化与空间调整的不同组合' in new
    # 旧轮次的逐字保持约束不能继续约束本轮审稿修订；保留数据与实物核验。
    revision_dir=ROOT/'outputs/审稿问题复核'
    revision=json.loads((revision_dir/'审稿复核结果.json').read_text(encoding='utf-8'))
    assert revision['master_sha256']==digest(ROOT/'outputs/沙洲观测主表.csv')
    grids=pd.read_csv(revision_dir/'Sentinel2原始栅格网格核验.csv')
    assert len(grids)==66 and grids.n_grids.eq(1).sum()==62
    assert set(grids.loc[grids.n_grids.gt(1),'sand_cay_id'])=={'GBR_037_cay_001','GBR_068_cay_001','GBR_095_cay_001','NH_015_cay_001'}
    position=pd.read_csv(revision_dir/'原始网格一致长期位置描述.csv')
    assert len(position)==90
    pd.testing.assert_frame_equal(position,pd.read_csv(src/'主图2_连续空间轨迹.csv'))
    forward=pd.read_csv(revision_dir/'仅过去训练年度预测逐景对照.csv')
    assert forward.groupby('metric').size().eq(1029).all()
    assert forward.sand_cay_id.nunique()==56
    assert (pd.to_datetime(forward.training_last_date).dt.year<forward.target_year).all()
    decomposition=pd.read_csv(revision_dir/'接受状态对绿色比例分解.csv')
    assert len(decomposition)==1490
    assert np.allclose(decomposition.fraction_change,decomposition.green_area_contribution+decomposition.denominator_contribution)
    changes=decomposition[decomposition.state_t0.ne(decomposition.state_t1)]
    assert len(changes)==172 and (changes.fraction_change*changes.green_area_change_m2<0).sum()==9
    assert (changes.denominator_contribution.abs()>changes.green_area_contribution.abs()).sum()==20
    assert '用于位置比较的90条长期观测序列' in new and '配对秩与平均增益的证据不一致' in new
    assert '53条Sentinel-2长期观测序列' in new and '−0.460' in new
    wave_current=pd.read_csv(revision_dir/'单一原始网格高波浪日模型.csv')
    wave_current=wave_current[wave_current.analysis.eq('E4_event_storm_days_bare')].iloc[0]
    assert '760个观测时段' in new and f'{int(wave_current.n_observations)}个裸沙时段' in new
    with zipfile.ZipFile(ROOT.parent/'归档/审稿主要问题修改前.zip') as before_review:
        verify_reference_revision(before_review.read('论文初稿.md').decode('utf-8'),new)
        current_doc=Document(ROOT/'论文/中文论文.docx')
        assert current_doc.paragraphs[0].text==new.splitlines()[0].removeprefix('# ')
        assert current_doc.core_properties.title==current_doc.paragraphs[0].text
        # 结果正文与源稿逐段对应，六幅图均位于各小节完整正文之后。
        paragraphs=current_doc.paragraphs
        first=next(i for i,p in enumerate(paragraphs) if p.text=='3 结果')
        last=next(i for i,p in enumerate(paragraphs) if p.text=='4 讨论')
        actual=[p.text for p in paragraphs[first:last] if p.text and not re.match(r'图 [1-6]｜',p.text)]
        section=new.split('## 3 结果',1)[1].split('## 4 讨论',1)[0]
        expected=['3 结果']+[block.removeprefix('### ').strip() for block in section.strip().split('\n\n') if block.strip()]
        assert actual==expected, 'Word结果正文与源稿不一致'
        # 讨论主文与源稿逐段对应，技术筛查明细继续由补充材料承接。
        discussion_end=next(i for i,p in enumerate(paragraphs) if p.text=='5 结论')
        actual_discussion=[p.text for p in paragraphs[last:discussion_end] if p.text]
        discussion_source=new.split('## 4 讨论',1)[1].split('## 5 结论',1)[0]
        expected_discussion=['4 讨论']+[block.removeprefix('### ').strip() for block in discussion_source.strip().split('\n\n') if block.strip()]
        assert actual_discussion==expected_discussion, 'Word讨论正文与源稿不一致'
        for number in range(1,7):
            begin=next(i for i,p in enumerate(paragraphs) if p.text.startswith(f'3.{number} '))
            finish=next(i for i,p in enumerate(paragraphs[begin+1:],begin+1) if p.text.startswith(f'3.{number+1} ') or p.text=='4 讨论')
            group=paragraphs[begin+1:finish]
            image_positions=[i for i,p in enumerate(group) if p._p.xpath('.//w:drawing')]
            assert len(image_positions)==1
            assert all(not p.text or p.text.startswith(f'图 {number}｜') for p in group[image_positions[0]+1:]), '图件打断小节正文'
            caption=next(p.text for p in group if p.text.startswith(f'图 {number}｜'))
            source_caption=next(line.replace('**','') for line in new.splitlines() if line.startswith(f'**图 {number}｜'))
            assert caption==source_caption, 'Word图注与源稿不一致'
        reference_heading=next(i for i,p in enumerate(current_doc.paragraphs) if p.text=='参考文献')
        word_references=[p.text.strip() for p in current_doc.paragraphs[reference_heading+1:] if p.text.strip()]
        expected_references=Counter(e.replace('*','') for e in reference_entries(new))
        assert Counter(word_references)==expected_references, 'Word参考文献内容或独立分段与源稿不一致'
    joint=pd.read_csv(ROOT/'额外分析_空间时间覆盖联合规律/八项联合检验.csv')
    assert len(joint)==8 and (joint.q<.05).sum()==1
    phase=pd.read_csv(ROOT/'outputs/季节相位质心分离统计.csv').set_index('bucket')
    assert np.isclose(phase.loc['same_season_cross_year','median_normalized_centroid_separation'],.126)
    assert np.isclose(phase.loc['opposite_season_within_year','median_normalized_centroid_separation'],.144)
    assert np.isclose(phase.loc['opposite_season_cross_year','median_normalized_centroid_separation'],.173)
    assert '补充表 S5-4' not in new and '结果见补充材料 S1–S3' not in new
    assert '覆盖全序列掩膜的固定裁剪窗' not in new
    results={'状态':'通过','原件指纹保持':True,'参考文献按核验清单修正':True,'文档实际核验':checked,'现行图数':len(manifest['图件']),'图件各格式核验数':len(manifest['图件'])*4,'表A':len(a),'表B':len(b),'几何缺测对':28,'ENSO两套结果表系数一致':True,'引用核验范围':'全文作者—年份对应及独立分段；59个原DOI与SAM修正DOI已获取登记元数据；全文逐句支持核查未全部完成'}
    results['连续空间及转换核查']={'连续序列':len(continuous),'可比首末':len(endpoints),'候选窗口':len(events),'不复用窗口':len(retained),'沙洲数':14,'严格四日期':1,'45天季节匹配':0,'NH背景筛查通过':9,'NH相邻区间':29,'解释':'描述性证据；未新增离散类型或控制后的转换响应模型'}
    results.update(现行图数=len(manifest['图件']),图件各格式核验数=4*len(manifest['图件']))
    results['伸长率补充核验']={'主序列':93,'观测':2115,'不重叠窗口':90,'边界情景保持方向':53,'完全45天同季匹配':0,'真实轮廓景数':10,'六幅主图结构保持':True,'新增S8':True}
    results['同季度对照']={'序列':63,'实际影像锚点':315,'边界方向保持':30,'与原窗口同符号':38,'两窗保持方向且同向':18,'原表与图源完全一致':True,'窗口位移已逐景复算':True}
    results['P0及主线恢复']={'紧致度同对象Crofton':True,'旧观测列保持':True,'紧致度年度增益':float(row.mean_relative_error_reduction),'紧致度年度q':float(row.fdr_q_value),'固定端点采样序列':93,'背景案例':10,'同区域配对':32,'不复用配对':13,'网格一致覆盖负关联保留':True,'表面状态图原影像选择保持且箭头已增强':True}
    results['3.1区域长期汇总']={'主序列':93,'重点四区':91,'首末交换可比':90,'面积分箱':[63,17,13],'源表逐值复核':True,'输入指纹一致':True,'段落数':len(new.split('### 3.1 ',1)[1].split('### 3.2 ',1)[0].strip().split('\n\n'))-1,'新增补充表':['S1-3','S1-4','S1-5']}
    results['章节职责与立意审校']={'标题与权威源稿一致':True,'摘要按审稿证据修订':True,'六主图八补图索引完整':True,'原始统计及阴性结果保留':True,'同季质心配对中位数':[.126,.144,.173],'旧补充编号修正':True,'方法裁剪窗说明与实现一致':True,'关键文献核验范围':'Kench2024、Lazarus2025/2026；2006官方摘要部分支持，不声称全部文献外部核验'}
    results['审稿主要问题复核']={'原始网格沙洲':66,'单一网格沙洲':62,'当前长期位置序列':90,'前向测试日期':1029,'前向沙洲':56,'状态对分解':1490,'独立配准精度':'未验证','盲重标':'未完成','公开本轮数据':('GitHub '+json.loads((ROOT/'outputs/公开归档核验.json').read_text(encoding='utf-8'))['version'] if (ROOT/'outputs/公开归档核验.json').exists() else '未发布'),'测量主表保持及参考文献按核验清单核对':True}
    from 复核审稿主要问题 import check_derived
    check_derived()
    sectors=pd.read_csv(ROOT/'outputs/植被扇区对照.csv')
    sector_stats=pd.read_csv(ROOT/'outputs/植被扇区对照统计.csv').set_index('stratum')
    background=pd.read_csv(revision_dir/'关键案例背景测量筛查.csv')
    accepted=background[background.specification.eq('physical_scale') & background.screen_accepted]
    assert '1012' in new and '994' in new and '平均低0.180' in new and '−0.330至−0.071' in new
    assert '没有直接检验各指标年度增益之间的差异' in new
    assert '36/167' in (ROOT/'论文补充材料.md').read_text(encoding='utf-8') and len(accepted)==36
    # 校验补充材料中的历史样本标识、交叉引用及其Word同步，避免与现行主模型混用。
    si_text=(ROOT/'论文补充材料.md').read_text(encoding='utf-8')
    si_word=Document(ROOT/'论文/中文论文补充材料.docx')
    si_paragraphs={p.text for p in si_word.paragraphs}
    for label in ['**表 S7-2｜','**表 S7-3｜','**表 S7-4｜']:
        line=next(s for s in si_text.splitlines() if s.startswith(label))
        assert '筛选前' in line and line.replace('**','') in si_paragraphs, '历史风浪样本标识或Word同步错误'
    current_wave=next(s for s in si_text.splitlines() if s.startswith('沿用高波浪日阈值及既有模型形式'))
    assert '季节控制见表S7-11' in current_wave and '逐区域剔除见表S7-12' in current_wave
    assert current_wave in si_paragraphs and f'{int(wave_current.n_observations)}个裸沙时段、{int(wave_current.n_cays)}个沙洲' in current_wave
    assert '正文主模型的1228' not in si_text and '主系数只在主图4a展示' not in si_text
    assert '留一年预测' not in si_text
    results['正文补充同步复核']={'当前高波浪时段':int(wave_current.n_observations),'当前高波浪沙洲':int(wave_current.n_cays),
        '历史样本单独标识':True,'季节与区域检验表号对应':True,'留一年重建与前向预测措辞区分':True}
    results['定向完善复核']={'图6b区间':len(sectors),'图6b沙洲':sectors.sand_cay_id.nunique(),
        '平均差':float(sector_stats.loc['all','mean_mobility_difference']),
        '聚类区间':[float(sector_stats.loc['all','bootstrap_ci_low']),float(sector_stats.loc['all','bootstrap_ci_high'])],
        '背景筛查通过':len(accepted),'背景筛查总区间':int(len(background)/2),
        '候选修正中位数米':float(accepted.background_correction_m.median()),
        '统计单位和年度措辞已修正':True,'独立测量缺口仍保留':True}
    # 方法正文逐段核对；公式迁移到补充后仍为可编辑数学对象。
    from 生成中文论文 import parse_markdown,clean_inline
    main_doc=Document(ROOT/'论文/中文论文.docx')
    start=next(i for i,p in enumerate(main_doc.paragraphs) if p.text=='2 方法')
    end=next(i for i,p in enumerate(main_doc.paragraphs) if p.text=='3 结果')
    word_methods=[p.text for p in main_doc.paragraphs[start:end] if p.text]
    source_blocks=parse_markdown(ROOT/'论文初稿.md')
    in_methods=False;expected_methods=[]
    for kind,content in source_blocks:
        if kind=='h1' and content=='2 方法':in_methods=True
        if kind=='h1' and content=='3 结果':break
        if in_methods and kind!='equation':expected_methods.append(clean_inline(content))
    assert word_methods==expected_methods, '方法源稿与Word逐段不一致'
    method_section=new.split('## 2 方法',1)[1].split('## 3 结果',1)[0]
    assert len(re.findall(r'^### 2\.',method_section,re.M))==7
    for required in ['Spearman','每个来源的六项','每个响应的12项','第90百分位','按礁盘联合重抽样','保留内部孔洞']:
        assert required in method_section,required
    for forbidden in ['填补内部孔洞','西北与东南控制点','主序列','原始网格']:
        assert forbidden not in method_section,forbidden
    with zipfile.ZipFile(ROOT.parent/'归档/方法重组与覆盖分组校正前.zip') as baseline:
        unchanged=['outputs/沙洲观测主表.csv','outputs/审稿问题复核/单一原始网格方向模型.csv',
                   'outputs/审稿问题复核/单一原始网格覆盖路径秩相关.csv','outputs/植被扇区对照统计.csv',
                   'outputs/审稿问题复核/单一原始网格留一年逐景对照.csv',
                   'outputs/审稿问题复核/单一原始网格留一年汇总.csv']
        for path in unchanged:
            assert baseline.read(path)==(ROOT/path).read_bytes(),path
    from lxml import etree
    equations=[]
    for name in ['中文论文.docx','中文论文补充材料.docx']:
        with zipfile.ZipFile(ROOT/'论文'/name) as z:
            tree=etree.fromstring(z.read('word/document.xml'))
            equations.append(len(tree.xpath('//m:oMath',namespaces={'m':'http://schemas.openxmlformats.org/officeDocument/2006/math'})))
    assert equations==[1,1]
    # 新增補充内容的段落、表格及公式均与Markdown逐项对应；图件插入不影响顺序。
    source_si=parse_markdown(ROOT/'论文补充材料.md')
    si_blocks=[]
    for kind,content in source_si:
        if kind=='h1' and content=='补充图注':break
        if kind=='equation':continue
        if kind=='table':si_blocks.append(('table',[[clean_inline(c) for c in row] for row in content]))
        else:si_blocks.append(('text',clean_inline(content)))
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    actual_si=[]
    for node in si_word._element.body:
        if node.tag==qn('w:tbl'):
            table=Table(node,si_word)
            actual_si.append(('table',[[c.text for c in row.cells] for row in table.rows]))
        elif node.tag==qn('w:p'):
            value=Paragraph(node,si_word).text
            if value and not value.startswith('补充图 '):actual_si.append(('text',value))
    assert actual_si==si_blocks,'补充材料源稿与Word段落或表格不同步'
    results['方法与覆盖分组核查']={'方法小节':7,'逐段Word同步':True,'主文补充可编辑公式':equations,
        '覆盖分组由逐时段审校标签验证':True,'旧阶段定义单独保留':True,
        '年度方向覆盖路径岸段原数据保持':True,'高波浪校正q':float(wave_current.fdr_q_value)}
    # 交叉引用和摘要分母在当前源稿、Word和原始统计表间对应。
    for manuscript in ['论文初稿.md','论文补充材料.md']:
        value=(ROOT/manuscript).read_text(encoding='utf-8')
        assert not re.search(r'(?:主文|方法|见|参见)\s*2\.[89](?![0-9])',value)
    abstract=new.split('## 摘要',1)[1].split('**关键词',1)[0].strip()
    assert all(x in abstract for x in ['93条','63条','90条','60条'])
    assert abstract in [p.text for p in main_doc.paragraphs]
    assert '对应版本v2.0.0' in new and '该仓库尚未包含本轮分析' not in new
    assert '状态定义同主文 2.5' in si_text and '模型设定见主文 2.6' in si_text
    public=json.loads((ROOT/'outputs/公开归档核验.json').read_text(encoding='utf-8'))
    assert public['all_remote_git_blob_hashes_verified']
    assert digest(ROOT/'公开数据代码.zip')==public['asset_sha256']
    results['公开归档与交叉引用']={'GitHub版本':public['version'],'GitHub提交':public['commit_sha'],
        '公开文件数':public['public_file_count'],'版本链接':public['release_url'],
        '公开包与远端指纹一致':True,'摘要面积序列及近稳':[93,63],'摘要位置序列及近稳':[90,60],
        '旧2.8及2.9引用消除':True,'S1模型及状态定义引用修正':True,'DOI':'按作者决定不申请'}
    (ROOT/'排版核验/交付核查.json').write_text(json.dumps(results,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(results,ensure_ascii=False,indent=2))

if __name__ == '__main__': main()
